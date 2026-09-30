from __future__ import annotations

import csv
import io
import json

from fastapi import APIRouter, Depends, HTTPException, Query
from fastapi.responses import StreamingResponse
from sqlalchemy import select
from sqlalchemy.ext.asyncio import AsyncSession

from app.db import get_session
from app.models import Record, RecordEvidence, Run, Source, Task
from app.schemas import TaskSpec

router = APIRouter(tags=["exports"])


def _sanitize_cell(value: str) -> str:
    """Neutralize formula injection in CSV/XLSX cells (including whitespace/tabs)."""
    if isinstance(value, str):
        stripped = value.lstrip(" \t\r\n")
        if stripped and stripped[0] in ("=", "+", "-", "@", "\t", "\r", "%"):
            return "'" + value
    return value


def _sanitize_row(row: dict) -> dict:
    """Sanitize all string values in a row dict."""
    return {k: _sanitize_cell(str(v)) if isinstance(v, str) else v for k, v in row.items()}


async def _load_records_with_provenance(
    run_id: str,
    field_names: list[str],
    session: AsyncSession,
) -> list[dict]:
    """Load records with provenance data in 3 batch queries instead of 1,000+ sequential queries."""
    records_result = await session.execute(
        select(Record).where(Record.run_id == run_id).order_by(Record.confidence.desc())
    )
    records = records_result.scalars().all()
    if not records:
        return []

    record_ids = [r.id for r in records]

    # Batch 1: Load all evidence for these records in a single query
    evidence_result = await session.execute(
        select(RecordEvidence).where(RecordEvidence.record_id.in_(record_ids))
    )
    all_evidence = evidence_result.scalars().all()

    # Group evidence by record_id
    evidence_by_record: dict[str, list[RecordEvidence]] = {}
    source_ids: set[str] = set()
    for ev in all_evidence:
        evidence_by_record.setdefault(ev.record_id, []).append(ev)
        if ev.source_id:
            source_ids.add(ev.source_id)

    # Batch 2: Load all referenced sources in a single query
    source_map: dict[str, str] = {}
    if source_ids:
        sources_result = await session.execute(
            select(Source.id, Source.url).where(Source.id.in_(list(source_ids)))
        )
        for s_id, s_url in sources_result.all():
            source_map[s_id] = s_url

    rows = []
    for r in records:
        data = json.loads(r.data)
        row = {name: data.get(name, "") for name in field_names}
        row["confidence"] = round(r.confidence, 3)

        ev_list = evidence_by_record.get(r.id, [])
        source_urls = [source_map[ev.source_id] for ev in ev_list if ev.source_id in source_map]
        snippets = [ev.snippet for ev in ev_list if ev.snippet]

        row["source_urls"] = " | ".join(source_urls)
        row["evidence"] = snippets[0] if snippets else ""
        row["retrieved_at"] = r.created_at.isoformat() if r.created_at else ""
        rows.append(row)

    return rows


@router.get("/runs/{run_id}/export")
async def export_records(
    run_id: str,
    format: str = Query(default="csv", pattern="^(csv|json|xlsx)$"),
    session: AsyncSession = Depends(get_session),
) -> StreamingResponse:
    run_result = await session.execute(select(Run).where(Run.id == run_id))
    run = run_result.scalar_one_or_none()
    if run is None:
        raise HTTPException(status_code=404, detail="Run not found")

    task_result = await session.execute(select(Task).where(Task.id == run.task_id))
    task = task_result.scalar_one_or_none()
    if task is None:
        raise HTTPException(status_code=404, detail="Task not found")

    # Use run_spec if available for column stability
    spec_json = run.run_spec if run.run_spec and run.run_spec != "{}" else task.spec
    spec = TaskSpec.model_validate_json(spec_json)
    field_names = [f.name for f in spec.fields]

    rows = await _load_records_with_provenance(run_id, field_names, session)

    if format == "json":
        # JSON includes nested evidence array per record
        json_rows = []
        for row in rows:
            record = {
                k: v for k, v in row.items() if k not in ("source_urls", "evidence", "retrieved_at")
            }
            record["confidence"] = row["confidence"]
            record["retrieved_at"] = row["retrieved_at"]

            # Parse source_urls back
            urls = [u.strip() for u in row.get("source_urls", "").split("|") if u.strip()]
            evidence_text = row.get("evidence", "")
            record["evidence"] = (
                [{"source_url": urls[0] if urls else None, "snippet": evidence_text}]
                if evidence_text
                else []
            )

            json_rows.append(record)

        content = json.dumps(json_rows, indent=2, ensure_ascii=False)
        return StreamingResponse(
            io.BytesIO(content.encode("utf-8")),
            media_type="application/json",
            headers={
                "Content-Disposition": f"attachment; filename=datalens_export_{run_id[:8]}.json"
            },
        )

    if format == "xlsx":
        try:
            from openpyxl import Workbook

            wb = Workbook()

            # Sheet 1: Records with provenance
            ws = wb.active
            ws.title = "Records"
            headers = field_names + ["confidence", "source_urls", "evidence", "retrieved_at"]
            ws.append(headers)
            for row in rows:
                sanitized = _sanitize_row(row)
                ws.append([sanitized.get(h, "") for h in headers])

            # Sheet 2: Sources
            sources_result = await session.execute(
                select(Source).where(Source.run_id == run_id).order_by(Source.fetched_at.desc())
            )
            sources = sources_result.scalars().all()
            ws2 = wb.create_sheet("Sources")
            ws2.append(
                ["url", "domain", "status", "http_status", "reason", "records_found", "fetched_at"]
            )
            for s in sources:
                ws2.append(
                    [
                        _sanitize_cell(s.url),
                        s.domain,
                        s.status.value,
                        s.http_status or "",
                        s.reason or "",
                        s.records_found,
                        s.fetched_at.isoformat() if s.fetched_at else "",
                    ]
                )

            # Sheet 3: Report (trust signals summary)
            stats = json.loads(run.stats) if run.stats else {}
            ws3 = wb.create_sheet("Report")
            ws3.append(["Metric", "Value"])
            ws3.append(["Raw extracted", stats.get("raw_count", 0)])
            ws3.append(["Verified", stats.get("verified_count", 0)])
            ws3.append(["Valid", stats.get("valid_count", 0)])
            ws3.append(["Deduplicated", stats.get("deduped_count", 0)])
            ws3.append(["Pages fetched", stats.get("pages_fetched", 0)])
            ws3.append(["Pages failed", stats.get("pages_failed", 0)])
            ws3.append(["Hallucinated dropped", stats.get("hallucinated_count", 0)])
            if stats.get("first_record_seconds"):
                ws3.append(["First record (seconds)", stats["first_record_seconds"]])
            if stats.get("duration_seconds"):
                ws3.append(["Total duration (seconds)", stats["duration_seconds"]])

            buf = io.BytesIO()
            wb.save(buf)
            buf.seek(0)
            return StreamingResponse(
                buf,
                media_type="application/vnd.openxmlformats-officedocument.spreadsheetml.sheet",
                headers={
                    "Content-Disposition": f"attachment; filename=datalens_export_{run_id[:8]}.xlsx"
                },
            )
        except ImportError:
            raise HTTPException(status_code=501, detail="XLSX export not available") from None

    # Default: CSV with provenance columns streamed via generator
    headers = field_names + ["confidence", "source_urls", "evidence", "retrieved_at"]

    def iter_csv():
        output = io.StringIO()
        writer = csv.DictWriter(output, fieldnames=headers, extrasaction="ignore")
        writer.writeheader()
        yield output.getvalue().encode("utf-8")
        output.seek(0)
        output.truncate(0)

        for row in rows:
            writer.writerow(_sanitize_row(row))
            yield output.getvalue().encode("utf-8")
            output.seek(0)
            output.truncate(0)

    return StreamingResponse(
        iter_csv(),
        media_type="text/csv",
        headers={"Content-Disposition": f"attachment; filename=datalens_export_{run_id[:8]}.csv"},
    )
